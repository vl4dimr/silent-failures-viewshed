# -*- coding: utf-8 -*-
"""
Auditoria del manuscrito: formato, figuras, coherencia numerica, prosa y
directrices de Journal of Archaeological Method and Theory (Springer).

La misma disciplina del caso de estudio, aplicada a este articulo. Cada cifra
del texto que procede de un calculo se coteja contra el JSON del que sale, y se
comprueba ademas que no este tecleada en el generador (p04). Las figuras y las
tablas deben estar citadas en orden; las referencias listadas, citadas por
apellido y ano, en cursiva donde APA lo pide y en orden alfabetico; las siglas,
definidas en su primera mencion; y la tipografia, la de la politica acordada
(sin rayas en el cuerpo, apostrofo y signo menos tipograficos, «p value»).

Salida: results/auditoria.json
"""
import json
import os
import re
import sys
import unicodedata
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.table import Table
from docx.text.paragraph import Paragraph

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RES = os.path.join(BASE, "results")
FIG = os.path.join(RES, "figuras")
CT = os.path.join(BASE, "data", "caso_titicaca")
DOCX = os.path.join(BASE, "manuscript_silent_failures.docx")
P04 = os.path.join(BASE, "scripts", "p04_manuscript.py")

ok_n, fail_n = 0, 0
items = []


def check(bloque, etiqueta, cond, detalle=""):
    global ok_n, fail_n
    if cond:
        ok_n += 1
    else:
        fail_n += 1
    items.append({"bloque": bloque, "check": etiqueta, "ok": bool(cond), "detalle": detalle})
    print("  %-5s %-62s %s" % ("OK" if cond else "FALLA", etiqueta, detalle), flush=True)


# formato de cifras: el mismo que p04 (signo menos tipografico)
def _m(s):
    return s.replace("-", "−")


def f(x, dec=2):
    return _m(("%%.%df" % dec) % x)


def fz(x):
    return _m("%+.2f" % x)


def pct(x, dec=0):
    return _m(("%%.%df" % dec) % (100.0 * x))


def g(x):
    return _m("%g" % x)


def num(n):
    return ("zero one two three four five six seven eight nine".split()[n]
            if 0 <= n <= 9 else str(n))


def bloques(doc):
    """Parrafos y tablas en el orden del documento."""
    for hijo in doc.element.body.iterchildren():
        if hijo.tag.endswith("}p"):
            yield Paragraph(hijo, doc)
        elif hijo.tag.endswith("}tbl"):
            yield Table(hijo, doc)


def main():
    doc = Document(DOCX)
    ps = [p.text for p in doc.paragraphs]
    texto = "\n".join(ps)
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                texto += "\n" + c.text

    # texto en orden de lectura (parrafos y celdas de tabla intercaladas)
    orden = []
    for b in bloques(doc):
        if isinstance(b, Paragraph):
            orden.append(b.text)
        else:
            for r in b.rows:
                orden.append(" | ".join(c.text for c in r.cells))
    i_ini = orden.index("1 Introduction")
    i_dec = orden.index("Statements and Declarations")
    i_refs = orden.index("References")
    cuerpo_orden = "\n".join(orden[i_ini:i_refs])        # cuerpo + declaraciones
    cuerpo_estricto = "\n".join(orden[i_ini:i_dec])       # sin declaraciones

    BEN = json.load(open(os.path.join(RES, "benchmark.json"), encoding="utf-8"))
    CAL = json.load(open(os.path.join(RES, "calibracion.json"), encoding="utf-8"))
    MOT = json.load(open(os.path.join(RES, "comparacion_motores.json"), encoding="utf-8"))
    FMETA = json.load(open(os.path.join(FIG, "meta.json"), encoding="utf-8"))
    fuente = open(P04, encoding="utf-8").read()

    print("AUDITORÍA DEL MANUSCRITO\n")

    # ---------------------------------------------------------------- formato
    print("1. FORMATO")
    runs = [r for p in doc.paragraphs for r in p.runs if r.text.strip()]
    fuentes = {r.font.name for r in runs}
    check("formato", "fuente única Arial", fuentes <= {"Arial"},
          "fuentes %s" % sorted(x for x in fuentes if x))
    sec = doc.sections[0]
    check("formato", "página A4 con márgenes de 2.5 cm",
          abs(sec.page_width - 7560000) < 40000 and abs(sec.top_margin / 360000 - 2.5) < 0.05,
          "%.1f x %.1f cm" % (sec.page_width / 360000, sec.page_height / 360000))
    i_p_ini = ps.index("1 Introduction")
    i_p_dec = ps.index("Statements and Declarations")
    cuerpo = [p for p in doc.paragraphs[i_p_ini:i_p_dec]
              if len(p.text.split()) > 25 and not re.match(r"^(Fig\.|Table)\s+\d+", p.text.strip())]
    just = sum(1 for p in cuerpo if p.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY)
    check("formato", "texto corrido justificado", just == len(cuerpo), "%d de %d" % (just, len(cuerpo)))
    presentes = [w for w in ("Mamani Calisaya", "Universidad Nacional del Altiplano",
                             "mmamanic@unap.edu.pe", "ORCID") if w in texto]
    check("formato", "bloque de autoría completo (anónimo simple)", len(presentes) == 4,
          "presentes: %d de 4" % len(presentes))
    corchetes = [p for p in ps if p.strip().startswith("[") and p.strip().endswith("]")]
    check("formato", "sin marcadores de posición entre corchetes", not corchetes)
    encab = [t.strip() for t in ps if re.match(r"^\d+(\.\d+)*\.?\s+[A-Z][a-z]", t.strip())
             and len(t.strip()) < 90]
    check("formato", "encabezados numerados sin punto final (plantilla Springer)",
          encab and not [e for e in encab if re.match(r"^\d+(\.\d+)*\.\s", e)],
          "%d encabezados" % len(encab))

    # ---------------------------------------------------------------- figuras
    print("\n2. FIGURAS Y TABLAS")
    from PIL import Image
    for fn in ("fig_pipeline.png", "fig_anatomia.png", "fig_laboratorio.png",
               "fig_consecuencias.png"):
        p = os.path.join(FIG, fn)
        im = Image.open(p)
        dpi = im.info.get("dpi", (0, 0))[0]
        check("figuras", "%s: resolución" % fn, dpi >= 299 and im.width >= 1800,
              "%dx%d px, %d dpi" % (im.width, im.height, dpi))
        check("figuras", "%s: existe en PDF vectorial" % fn,
              os.path.exists(p.replace(".png", ".pdf")))
    check("figuras", "cuatro figuras incrustadas", len(doc.inline_shapes) == 4,
          "%d" % len(doc.inline_shapes))
    for i, s in enumerate(doc.inline_shapes, 1):
        w, h = s.width / 36000, s.height / 36000
        check("figuras", "Fig. %d incrustada dentro de 174 x 234 mm" % i, w <= 174 and h <= 234,
              "%.1f x %.1f mm" % (w, h))
    pies = [t.strip() for t in ps if re.match(r"^Fig\.\s+\d+\s", t.strip())]
    check("figuras", "cuatro pies «Fig. n»", len(pies) == 4, "%d" % len(pies))
    check("figuras", "pies de figura sin puntuación final",
          not [p for p in pies if p[-1] in ".;:"])
    colores = re.compile(r"\b(red|blue|green|grey|gray|orange|yellow|purple|black|colou?r)\b", re.I)
    check("figuras", "pies de figura sin referencias al color (impresión en B/N)",
          not [p for p in pies if colores.search(p)])

    # rotulos de figura y tabla: parrafos cuya primera run es «Fig. n»/«Table n» en negrita
    rotulos = {p.text for p in doc.paragraphs
               if p.runs and p.runs[0].bold
               and re.match(r"^(Fig\.|Table)\s+\d+\s*$", p.runs[0].text.strip())}

    def primeras(patron, excluir):
        vistos = []
        for t in orden[i_ini:i_refs]:
            if t in rotulos:
                continue
            for m in re.finditer(patron, t):
                if m.group(1) not in vistos:
                    vistos.append(m.group(1))
        return vistos

    of = primeras(r"Fig\.\s*(\d)", r"^Fig\.\s+\d+\s")
    check("figuras", "figuras citadas por primera vez en orden 1, 2, 3, 4", of == ["1", "2", "3", "4"],
          "orden: %s" % of)
    n_tab = len(doc.tables)
    ot = primeras(r"Table\s+(\d)", r"^Table\s+\d+\s")
    check("figuras", "tablas citadas por primera vez en orden consecutivo",
          ot == [str(i) for i in range(1, n_tab + 1)], "orden: %s" % ot)
    rot = [t for t in rotulos if t.startswith("Table")]
    check("figuras", "un rótulo «Table n» por tabla", len(rot) == n_tab, "%d / %d" % (len(rot), n_tab))

    # ------------------------------------------------------------- coherencia
    print("\n3. COHERENCIA NUMÉRICA (cada cifra, contra su JSON)")
    MC = BEN["motor_correcto"]
    CASOS = MC["casos"]
    N_TESTS, N_CASOS, N_PROPS = MC["pruebas"], len(CASOS), len(MC["propiedades"])
    check("números", "pruebas, casos y propiedades del banco",
          ("all %d tests (%d cases and %s properties)" % (N_TESTS, N_CASOS, num(N_PROPS))) in texto
          and ("%d terrain cases" % N_CASOS) in texto
          and ("%d behavioural properties" % N_PROPS) in texto
          and ("consists of %s behavioural properties" % num(N_PROPS)) in texto,
          "%d = %d + %d" % (N_TESTS, N_CASOS, N_PROPS))
    RG = MC["comprobacion_rasgo_decide"]
    check("números", "casos decisivos y controles",
          RG["ok"] and ("Of the %d cases, %d are feature cases and %d are controls"
                        % (N_CASOS, RG["casos_de_rasgo"], RG["controles"])) in texto,
          "%d + %d" % (RG["casos_de_rasgo"], RG["controles"]))
    n_ev = sum(p["comprobaciones"] for p in MC["propiedades"])
    check("números", "evaluaciones de las propiedades", ("{:,} evaluations".format(n_ev)) in texto,
          "{:,}".format(n_ev))
    DC = MC["distancia_critica"]
    check("números", "distancia crítica exacta y la aproximada de la v1",
          ("gives %s km" % f(DC["exacta_m"] / 1000, 1)) in texto
          and ("gave %s km instead of %s km" % (f(DC["aproximada_v1_m"] / 1000, 1),
                                               f(DC["exacta_m"] / 1000, 1))) in texto,
          "%s / %s km" % (f(DC["exacta_m"] / 1000, 1), f(DC["aproximada_v1_m"] / 1000, 1)))
    check("números", "celda del banco", ("grid of %s m cells" % g(MC["celda_m"])) in texto,
          "%s m" % g(MC["celda_m"]))

    # tabla de los casos
    tcasos = next(t for t in doc.tables if t.rows[0].cells[0].text.startswith("Case"))
    filas = [[c.text for c in r.cells] for r in tcasos.rows[1:]]
    bien = len(filas) == N_CASOS and all(
        fl[2] == f(c["distancia_km"], 2) and fl[3] == "%s / %s" % (g(c["h_obs"]), g(c["h_tgt"]))
        and fl[4] == ("feature case" if c["tipo"] == "rasgo" else "control")
        for fl, c in zip(filas, CASOS))
    check("números", "tabla de casos: distancia, alturas y papel de cada caso", bien,
          "%d filas" % len(filas))

    MUT = BEN["mutantes"]
    n_comp = MUT["extremos_incluidos"]["comparaciones_equivalencia"]
    check("números", "barrido de equivalencia", ("{:,}".format(n_comp)) in texto, "{:,}".format(n_comp))
    ce = MUT["extremos_incluidos"]["contraejemplo_alturas_negativas"]
    check("números", "contraejemplo con altura negativa",
          ("%s m below the ground" % g(-ce["h_obs"])) in texto
          and ("%s km away" % f(ce["distancia_km"], 2)) in texto,
          "h_obs = %s" % g(ce["h_obs"]))
    PV1 = BEN["propiedades_v1"]
    tmut = next(t for t in doc.tables if t.rows[0].cells[0].text == "Engine candidate")
    fm = {r.cells[0].text: [c.text for c in r.cells] for r in tmut.rows[1:]}
    NOMBRE = {"curvatura_restada": "Curvature subtracted", "sin_curvatura": "Curvature omitted",
              "altura_objetivo_nula": "Target height zero", "extremos_incluidos": "Endpoints included",
              "muestreo_grueso": "Coarse profile sampling"}
    for d, v in MUT.items():
        fila = fm.get(NOMBRE[d], [])
        esper = ["%d/%d" % (v["casos_que_saltan"], v["de_casos"]),
                 "%d/%d" % (v["propiedades_que_saltan"], v["de_propiedades"]),
                 "%d/%d" % (len(PV1["cuales_por_mutante"][d]), len(PV1["nombres"]))]
        check("números", "matriz de mutantes: %s (casos, propiedades, v1)" % d,
              fila[1:4] == esper, " · ".join(esper))
    check("números", "propiedades v1: ningún mutante detectado",
          PV1["mutantes_detectados"] == 0
          and "properties detected none of them" in texto
          and "properties detected none of the injected defects" in texto,
          "%d de %d" % (PV1["mutantes_detectados"], PV1["mutantes_evaluados"]))
    check("números", "todo mutante no equivalente muere por caso y por propiedad",
          BEN["resumen"]["todo_mutante_no_equivalente_muere_por_caso_y_propiedad"])

    # casos trasladados por la comprobacion del rasgo
    def dist(nombre):
        return f(next(c for c in CASOS if c["prueba"] == nombre)["distancia_km"], 0)
    tras = (dist("barrera de 50 m a 6 km: tapa"), dist("la misma barrera de 2 m si tapa a 9 km"),
            dist("objetivo a ras de suelo tras loma de 12 m a 4 km: no se ve"))
    check("números", "distancias a las que se trasladaron los casos",
          ("moved to %s, %s and %s km" % tras) in texto, "%s, %s, %s km" % tras)

    RESU = CAL["resumen"]
    for esc in ("S0", "S1"):
        for proc in ("correcto", "agua", "recorte"):
            c = RESU[esc][proc]
            par = "%s (%s)" % (f(c["z_media"], 2), f(c["z_sd"], 2))
            check("números", "tabla de calibración: z de %s/%s" % (esc, proc), par in texto, par)
    c0 = RESU["S0"]["correcto"]
    ic = "%s %% (95 %% confidence interval, CI, %s–%s %%" % (
        pct(c0["rechazo"], 1), pct(c0["rechazo_ic95"][0], 1), pct(c0["rechazo_ic95"][1], 1))
    check("números", "rechazo bajo S0 e intervalo de Wilson", ic.replace("%%", "%") in texto, ic)
    check("números", "potencia teórica",
          ("%s %%" % pct(RESU["potencia_teorica_S1_correcto"])) in texto,
          pct(RESU["potencia_teorica_S1_correcto"]) + " %")
    m = RESU["mecanismo_recorte"]
    check("números", "rho de Spearman y desplazamiento medio del recorte",
          ("ρ = %s" % f(m["spearman_desalineacion"], 2)) in texto
          and ("procedure is %s:" % f(m["dz_medio"], 2)) in texto
          and ("standard deviation of %s" % f(m["dz_sd"], 2)) in texto and m["p_spearman"] < 0.001,
          f(m["spearman_desalineacion"], 2))
    dgn = RESU["diagnosticos"]
    check("números", "diagnósticos D1 y D2",
          ("%s %% of the null placements accepted" % pct(dgn["agua_nulos_que_pisan_media"])).replace("%%", "%") in texto
          and ("admitted on average %s %% of the" % pct(dgn["cobertura_recorte_S0_media"])).replace("%%", "%") in texto,
          "%s %% / %s %%" % (pct(dgn["agua_nulos_que_pisan_media"]), pct(dgn["cobertura_recorte_S0_media"])))
    CFG = CAL["config"]
    check("números", "réplicas y colocaciones nulas del laboratorio",
          ("In total, %d replicates were run, each with N = %d null placements"
           % (CFG["replicas"], CFG["n_null"])) in texto, "%d, %d" % (CFG["replicas"], CFG["n_null"]))
    check("números", "sitios de la nube sintética y orientaciones de D2",
          ("cloud of %d sites" % CFG["n_sitios"]) in texto
          and ("the %d placement orientations" % CFG["pasos_orientacion"]) in texto
          and ("all %d orientations" % CFG["pasos_orientacion"]) in texto,
          "%d sitios, %d orientaciones" % (CFG["n_sitios"], CFG["pasos_orientacion"]))
    ruido = 1.0 / (CFG["n_null"] ** 0.5)
    check("números", "ruido de Monte Carlo 1/√N", ("about %s for N = %d" % (f(ruido, 2), CFG["n_null"])) in texto,
          f(ruido, 2))

    # caso de campo
    def cj(fn):
        return json.load(open(os.path.join(CT, fn), encoding="utf-8"))
    for fn, eti in (("nulo_rigido.json", "z corregida del caso de campo"),
                    ("nulo_rigido_sin_mascara.json", "z del agua sin enmascarar"),
                    ("nulo_rigido_n300.json", "z del recorte del caso de campo")):
        v = fz(cj(fn)["por_alcance"]["5000"]["z"])
        check("números", eti, v in texto, v)
    agua = pct(cj("nulo_rigido.json")["fraccion_agua"], 1)
    agua_r = pct(cj("nulo_rigido_n300.json")["fraccion_agua"], 1)
    check("números", "fracción de agua del área corregida (contraste sin máscara)",
          ("covering %s %% of the study area" % agua).replace("%%", "%") in texto, agua + " %")
    check("números", "fracción de agua del recorte",
          ("lake covered %s %% of the" % agua_r).replace("%%", "%") in texto, agua_r + " %")
    n_campo = cj("nulos.json")["parametros"]["n_sitios"]
    check("números", "sitios del estudio de campo", ("study of %d archaeological sites" % n_campo) in texto,
          "%d" % n_campo)
    mo = re.search(r"«(\d+) de (\d+) orientaciones»",
                   open(os.path.join(CT, "README_origen.txt"), encoding="utf-8").read())
    check("números", "orientaciones del recorte, declaradas como cifra del estudio de campo",
          ("field study reports that a tightly fitted raster crop" in texto)
          and ("admitted only %s of the %s orientations" % mo.groups()) in texto
          and ("reports %s of %s for its own crop" % mo.groups()) in texto,
          "%s/%s (README_origen.txt)" % mo.groups())
    check("números", "distancia de la figura 2", ("%s km" % f(FMETA["fig1_km"], 1)) in texto,
          "%s km" % f(FMETA["fig1_km"], 1))

    # motores
    G = MOT["motores"]["GDAL gdal_viewshed"]
    R = MOT["motores"]["GRASS r.viewshed"]
    tmot = next(t for t in doc.tables if t.rows[0].cells[0].text == "Engine")
    fg = [c.text for c in tmot.rows[1].cells]
    fr = [c.text for c in tmot.rows[2].cells]
    check("números", "gdal_viewshed: configurado, de fábrica y dirección",
          fg[1:5] == [G["correcto"], G["de_fabrica"],
                      str(G["direccion"]["declara_tapado_lo_visible"]),
                      str(G["direccion"]["declara_visible_lo_tapado"])] and G["version_gdal"] in fg[0],
          " · ".join(fg[1:5]))
    check("números", "r.viewshed: configurado, de fábrica y dirección",
          fr[1:5] == [R["correcto"], R["de_fabrica"],
                      str(R["direccion"]["declara_tapado_lo_visible"]),
                      str(R["direccion"]["declara_visible_lo_tapado"])] and R["version"] in fr[0],
          " · ".join(fr[1:5]))
    atr_g = "target: %d/%d; observer: %d/%d; curvature: %d/%d" % tuple(
        x for k in ("tz", "oz", "cc") for x in (G["atribucion"][k]["coinciden"], G["atribucion"][k]["de"]))
    atr_r = "target: %d/%d; observer: %d/%d; curvature and refraction: %d/%d" % tuple(
        x for k in ("target_elevation", "observer_elevation", "curvatura_y_refraccion")
        for x in (R["atribucion"][k]["coinciden"], R["atribucion"][k]["de"]))
    check("números", "atribución de cada fallo a un valor de fábrica", fg[5] == atr_g and fr[5] == atr_r,
          "%s | %s" % (atr_g, atr_r))
    check("números", "texto: -tz sola recupera todos los casos de GDAL",
          ("restoring the target height alone recovers %d of %d cases"
           % (G["atribucion"]["tz"]["coinciden"], G["atribucion"]["tz"]["de"])) in texto)
    D = MOT["discrepancia_entre_motores_de_fabrica"]
    n_comp_err = MOT["coincidencia_en_el_error_de_fabrica"]["casos"]
    n_bien = MOT["ambos_aciertan_de_fabrica"]["casos"]
    check("números", "discrepancias y errores compartidos entre motores de fábrica",
          ("disagree on %s of the %d cases and agree on %d" % (num(D["casos"]), N_CASOS, N_CASOS - D["casos"])) in texto
          and ("%s share the same wrong verdict" % num(n_comp_err)) in texto
          and ("disagreeing on %d cases and returning the same wrong verdict on %d others"
               % (D["casos"], n_comp_err)) in texto
          and ("disagree on %d cases, return the same wrong verdict on %d and are both right on %d"
               % (D["casos"], n_comp_err, n_bien)) in texto
          and ("gdal_viewshed is right on %s of them and r.viewshed on %s"
               % (num(D["acierta_gdal"]), num(D["acierta_grass"]))) in texto,
          "discrepan %d, error común %d, ambos bien %d" % (D["casos"], n_comp_err, n_bien))
    cr = R["comprobacion_refraccion"]

    def umbral(modo):
        v = next(c["veredictos"] for c in cr["corridas"] if c["modo"] == modo)
        ks = cr["distancias_km"]
        return (max((k for k in ks if v[k]), key=float), min((k for k in ks if not v[k]), key=float))
    u1 = umbral("-c refraction_coeff=0.13 (sin -r; protocolo de la primera version)")
    u2 = umbral("-c -r refraction_coeff=0.13")
    check("números", "comprobación de la bandera -r de r.viewshed",
          ("lost between %s and %s km, exactly as with" % u1) in texto
          and ("it was lost between %s and %s km" % u2) in texto and cr["coef_sin_r_es_igual_que_c_solo"],
          "%s–%s frente a %s–%s km" % (u1 + u2))
    check("números", "versiones de GDAL y GRASS", ("version %s (GDAL/OGR" % G["version_gdal"]) in texto
          and ("version %s (GRASS Development" % R["version"]) in texto,
          "%s / %s" % (G["version_gdal"], R["version"]))

    # ninguna de estas cifras tecleada en el generador
    literales = {
        "distancia crítica": f(DC["exacta_m"] / 1000, 1) + " km",
        "fracción de agua": agua, "fracción de agua del recorte": agua_r,
        "sitios del campo": "%d archaeological" % n_campo,
        "sitios de la nube": "%d sites" % CFG["n_sitios"],
        "orientaciones": "%d orientations" % CFG["pasos_orientacion"],
        "orientaciones del recorte": mo.group(1),
        "barrido": "{:,}".format(n_comp), "barrido sin coma": str(n_comp),
        "configurado": G["correcto"], "GDAL de fábrica": G["de_fabrica"],
        "GRASS de fábrica": R["de_fabrica"],
        "versión de GDAL": G["version_gdal"], "versión de GRASS": R["version"],
        "rho": f(m["spearman_desalineacion"], 2).replace("−", "-"),
    }
    tecleadas = [k for k, v in literales.items() if v in fuente]
    check("números", "ninguna de esas cifras está tecleada en p04", not tecleadas,
          "tecleadas: %s" % tecleadas if tecleadas else "%d cifras vigiladas" % len(literales))

    # -------------------------------------------------------------------- prosa
    print("\n4. PROSA Y APARATO CRÍTICO")
    refs = [t.strip() for t in orden[i_refs + 1:] if t.strip()]
    ref_ps = [p for p in doc.paragraphs[ps.index("References") + 1:] if p.text.strip()]

    def autores_y_ano(r):
        m = re.match(r"^(.*?) \(((?:19|20)\d{2})[a-z]?\)\.", r)
        if not m:
            return None, [], None
        aut, ano = m.group(1), m.group(2)
        noms = re.findall(r"(?:^|, |& )([^,&]+?), (?:[A-Z]\.(?:[ -]?[A-Z]\.)*)", aut)
        if not noms:
            return aut.rstrip("."), [aut.rstrip(".")], ano
        return noms[0].strip(), [n.strip() for n in noms], ano

    sin_citar, mal_formadas = [], []
    for r in refs:
        ap, noms, ano = autores_y_ano(r)
        if not ap:
            mal_formadas.append(r[:40])
            continue
        a = re.escape(ap)
        if len(noms) == 1:
            pat = r"%s(?:’s)? \((?:\d{4}, )*%s|%s, %s(?!\d)" % (a, ano, a, ano)
        elif len(noms) == 2:
            b = re.escape(noms[1])
            pat = r"%s and %s \(%s|%s & %s, %s" % (a, b, ano, a, b, ano)
        else:
            pat = r"%s et al\.(?: \(|, )%s" % (a, ano)
        if not re.search(pat, cuerpo_orden):
            sin_citar.append("%s %s" % (ap, ano))
    check("prosa", "toda referencia listada está citada (apellido y año)",
          not sin_citar and not mal_formadas,
          ("sin citar: %s" % sin_citar) if sin_citar else
          ("mal formadas: %s" % mal_formadas) if mal_formadas else "%d referencias" % len(refs))

    def clave(r):
        aut = re.match(r"^(.*?) \((?:19|20)\d{2}", r).group(1)
        ano = re.search(r"\(((?:19|20)\d{2})", r).group(1)
        s = unicodedata.normalize("NFKD", aut)
        s = "".join(ch for ch in s if not unicodedata.combining(ch)).lower()
        return re.sub(r"[^a-z ]", "", s).split() + [ano]
    check("prosa", "lista de referencias en orden alfabético",
          [clave(r) for r in refs] == sorted(clave(r) for r in refs))
    sin_cursiva = [p.text[:30] for p in ref_ps if not any(r.italic for r in p.runs)]
    check("prosa", "toda referencia lleva en cursiva revista, libro o programa", not sin_cursiva,
          "sin cursiva: %s" % sin_cursiva[:3] if sin_cursiva else "%d" % len(ref_ps))
    check("prosa", "ningún asterisco de marcado llega al documento", "*" not in "\n".join(refs))
    con_doi = [r for r in refs if "https://doi.org/" in r]
    check("prosa", "toda referencia con DOI lo lleva como enlace completo",
          len(con_doi) >= len(refs) - 1, "%d de %d" % (len(con_doi), len(refs)))
    check("prosa", "depósito del estudio de campo citado en su versión con DOI de versión",
          "(Version 1.2.0)" in "\n".join(refs) and "10.5281/zenodo.23050689" in "\n".join(refs))
    check("prosa", "depósito propio en la lista y en la disponibilidad de datos",
          "zenodo.22242923" in "\n".join(refs) and "(Mamani Calisaya, 2026)" in cuerpo_orden)

    americanismos = [w for w in ("artifact", "behavior", "modeling", "analyzed", "signaling",
                                 "color", "center", "upward")
                     if re.search(r"\b%s\b" % w, cuerpo_orden)]
    check("prosa", "inglés británico consistente", not americanismos,
          "revisar: %s" % americanismos if americanismos else "")

    # politica tipografica
    check("prosa", "sin rayas (—) en el cuerpo", "—" not in cuerpo_orden,
          "%d" % cuerpo_orden.count("—"))
    rectos = [t[:40] for t in orden if "'" in t]
    check("prosa", "sin apóstrofos rectos", not rectos, "%s" % rectos[:2] if rectos else "")
    negativos = re.findall(r"(?<![\w‑.\-/])-\d", cuerpo_orden)
    check("prosa", "negativos con signo menos (U+2212)", not negativos, "%d" % len(negativos))
    prohibidas = ["p-value", "hand-written", "Hand-written", "factory", "failing in opposite directions",
                  "opposite directions", "7 of 15", "nearly half", "QGIS menu entry", "presses Run",
                  "QGIS processing framework", "through the QGIS processing", "reached through QGIS",
                  "same menu", "never saw", "commonly used", "Neither, they note",
                  "a decade of fiscal", "exceeded the signal", "the only route", "assessed probabilistic "
                  "visibilities against field", "consistent with the laboratory"]
    sin_refs = "\n".join(orden[:i_refs])
    halladas = [x for x in prohibidas if x in sin_refs]
    check("prosa", "sin frases retiradas por los informes", not halladas,
          "quedan: %s" % halladas if halladas else "%d vigiladas" % len(prohibidas))

    # siglas definidas en su primera mencion
    SIGLAS = {"GIS": "geographic information system", "DEM": "digital elevation model",
              "KS": "Kolmogorov", "CI": "confidence interval",
              "CDF": "cumulative distribution function",
              "GDAL": "Geospatial Data Abstraction Library",
              "GRASS": "Geographic Resources Analysis Support System",
              "QGIS": "Quantum GIS", "LLM": "large language model", "SD": "standard deviation"}
    PERMITIDAS = {"MIT", "OGR", "ORCID"}
    for s, exp in SIGLAS.items():
        mm = re.search(r"\b%ss?\b" % s, cuerpo_orden)
        if not mm:
            check("jamt", "sigla %s definida en su primera mención" % s, True, "no se usa")
            continue
        ventana = cuerpo_orden[max(0, mm.start() - 120):mm.end() + 120]
        check("jamt", "sigla %s definida en su primera mención" % s, exp.lower() in ventana.lower())
    otras = sorted(set(re.findall(r"\b[A-Z]{2,6}\b", cuerpo_orden)) - set(SIGLAS) - PERMITIDAS)
    check("jamt", "ninguna otra sigla sin definir en el cuerpo", not otras, "%s" % otras if otras else "")

    # ------------------------------------------------------ directrices de JAMT
    print("\n4b. DIRECTRICES DE JAMT (SPRINGER)")
    i_ab = ps.index("Abstract")
    i_kw = ps.index("Keywords")
    abstract = " ".join(ps[i_ab + 1:i_kw])
    n_abs = len(abstract.split())
    check("jamt", "resumen de 150 a 250 palabras", 150 <= n_abs <= 250, "%d" % n_abs)
    check("jamt", "resumen sin siglas", not re.findall(r"\b[A-Z]{2,6}\b", abstract),
          "%s" % re.findall(r"\b[A-Z]{2,6}\b", abstract))
    check("jamt", "resumen sin referencias", not re.search(r"\((?:19|20)\d{2}\)|, (?:19|20)\d{2}\)", abstract))
    n_kw = len([k for k in ps[i_kw + 1].split(";") if k.strip()])
    check("jamt", "de 4 a 6 palabras clave", 4 <= n_kw <= 6, "%d" % n_kw)
    palabras_total = sum(len(p.split()) for p in ps) + sum(
        len(c.text.split()) for t in doc.tables for r in t.rows for c in r.cells)
    check("jamt", "extensión de artículo de método (6 000 a 12 000 palabras)",
          6000 <= palabras_total <= 12000, "%d" % palabras_total)
    harvard = re.findall(r"\b([A-Z][A-Za-z'’-]+(?: et al\.)?) (19|20)\d{2}[;)]", cuerpo_orden)
    check("jamt", "citas APA, con coma antes del año", not harvard,
          "quedan: %s" % [h[0] for h in harvard[:3]] if harvard else "")
    and_par = re.findall(r"\([^()]*\b[A-Z][a-z]+ and [A-Z][a-z]+, (?:19|20)\d{2}", cuerpo_orden)
    check("jamt", "«&» dentro de las citas parentéticas", not and_par)
    sin_par = [r[:44] for r in refs if not re.search(r"\((19|20)\d{2}[a-z]?\)\.", r)]
    check("jamt", "referencias con el año entre paréntesis", not sin_par)
    niveles = {e.split()[0].count(".") for e in encab}
    check("jamt", "encabezados numerados de hasta tres niveles", niveles and max(niveles) <= 2,
          "profundidades: %s" % sorted(niveles))
    bloque = "\n".join(ps[ps.index("Statements and Declarations"):ps.index("References")])
    faltan = [k for k in ("Funding", "Competing interests", "Data and code availability",
                          "Author contributions") if k not in bloque]
    check("jamt", "Statements and Declarations completas antes de las referencias", not faltan)
    i_met = ps.index("3 Materials and methods")
    i_res = ps.index("4 Results")
    metodos = "\n".join(ps[i_met:i_res]).lower()
    check("jamt", "uso de modelo de lenguaje documentado en Métodos",
          "large language model" in metodos and "claude" in metodos)

    # ------------------------------------------------- propiedades del fichero
    print("\n5. PROPIEDADES DEL FICHERO")
    cp = doc.core_properties
    rastro = [k for k in ("author", "last_modified_by", "comments", "category", "subject", "keywords")
              if (getattr(cp, k, "") or "").strip()]
    check("fichero", "sin autor ni rastro del generador en los metadatos", not rastro)
    with zipfile.ZipFile(DOCX) as z:
        crudo = (z.read("docProps/core.xml") + z.read("docProps/app.xml")).decode("utf-8")
    sospechas = [w for w in ("python-docx", "Macintosh", "Alanoca") if w in crudo]
    check("fichero", "los metadatos XML no delatan nada", not sospechas)
    check("fichero", "fecha de creación verosímil", cp.created is not None and cp.created.year >= 2026)

    print("\n" + "=" * 70)
    print("RESULTADO: %d comprobaciones correctas, %d fallos" % (ok_n, fail_n))
    print("=" * 70)
    json.dump({"ok": ok_n, "fallos": fail_n, "items": items},
              open(os.path.join(RES, "auditoria.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=2)
    return 1 if fail_n else 0


if __name__ == "__main__":
    sys.exit(main())
