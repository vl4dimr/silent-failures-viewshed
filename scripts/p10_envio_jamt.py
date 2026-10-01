# -*- coding: utf-8 -*-
"""
Paquete de envio a Journal of Archaeological Method and Theory (Springer).

Springer pide el manuscrito en Word con la portada dentro (revision de anonimo
simple), las figuras como ficheros aparte nombrados «Fig1», «Fig2», … en un
formato de su lista (vectorial: EPS o PDF; tramas: TIFF) y una carta de
presentacion. Este guion copia el manuscrito ya generado y auditado, copia cada
figura en PDF vectorial y en TIFF (nunca el .svg, que no esta en la lista de
Springer) y redacta la carta con cifras leidas del propio manuscrito, para que
no se desvien de el.

Salida: ENVIO_JAMT/
    1_MANUSCRIPT.docx
    2_COVER_LETTER.docx  (+ .txt para pegar en Editorial Manager)
    Fig1.pdf … Fig4.pdf  y  Fig1.tif … Fig4.tif
"""
import os
import re
import shutil
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm
from PIL import Image

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docmeta

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIG = os.path.join(BASE, "results", "figuras")
MS = os.path.join(BASE, "manuscript_silent_failures.docx")
OUT = os.path.join(BASE, "ENVIO_JAMT")
os.makedirs(OUT, exist_ok=True)

AUTOR = "Milton Vladimir Mamani Calisaya"
FILIACION = "Universidad Nacional del Altiplano, Puno, Peru"
CORREO = "mmamanic@unap.edu.pe"
ORCID = "0000-0002-0676-0989"
DOI_ESTE = "10.5281/zenodo.22242923"

# ----------------------------------------------------------- cifras del manuscrito
src = Document(MS)
ps = [p.text.strip() for p in src.paragraphs]
titulo = ps[0] + ": " + ps[1][0].lower() + ps[1][1:]
i_ab, i_kw = ps.index("Abstract"), ps.index("Keywords")
resumen = " ".join(ps[i_ab + 1:i_kw])
n_abs = len(resumen.split())
i_ini = ps.index("1 Introduction")
i_dec = ps.index("Statements and Declarations")
i_ref = ps.index("References")
n_cuerpo = sum(len(t.split()) for t in ps[i_ini:i_dec])
n_tablas = sum(len(c.text.split()) for t in src.tables for r in t.rows for c in r.cells)
n_refs = len([t for t in ps[i_ref + 1:] if t])
n_fig = len(src.inline_shapes)
n_tab = len(src.tables)
n_casos = int(re.search(r"\((\d+) terrain cases", resumen).group(1))
n_disc, n_comun = (int(x) for x in re.search(
    r"disagreeing on (\d+) cases and returning the same wrong verdict on (\d+) others",
    resumen).groups())

# ------------------------------------------------------------------- copias
for fn in os.listdir(OUT):
    if re.match(r"^Fig\d\.", fn):          # restos de paquetes anteriores (png, svg)
        os.remove(os.path.join(OUT, fn))
shutil.copyfile(MS, os.path.join(OUT, "1_MANUSCRIPT.docx"))
for n, base in enumerate(["fig_pipeline", "fig_anatomia", "fig_laboratorio",
                          "fig_consecuencias"], start=1):
    shutil.copyfile(os.path.join(FIG, base + ".pdf"), os.path.join(OUT, "Fig%d.pdf" % n))
    # TIFF de produccion a 600 ppp, rasterizado desde el PDF vectorial: Springer
    # pide 600 ppp para figuras que combinan trazo y tono continuo.
    import fitz
    with fitz.open(os.path.join(FIG, base + ".pdf")) as pdf:
        pix = pdf[0].get_pixmap(dpi=600, alpha=False)
    im = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
    im.save(os.path.join(OUT, "Fig%d.tif" % n), dpi=(600, 600), compression="tiff_lzw")

# comprobaciones del paquete
ficheros = sorted(os.listdir(OUT))
assert not [f for f in ficheros if f.lower().endswith(".svg")], "queda un .svg"
for n in range(1, 5):
    assert "Fig%d.pdf" % n in ficheros and "Fig%d.tif" % n in ficheros, n
    im = Image.open(os.path.join(OUT, "Fig%d.tif" % n))
    assert im.mode == "RGB" and im.info.get("dpi", (0, 0))[0] >= 599, (n, im.mode, im.info.get("dpi"))
    assert im.width / 600 * 25.4 <= 174.5, ("ancho en mm", n, im.width / 600 * 25.4)
    with open(os.path.join(OUT, "Fig%d.pdf" % n), "rb") as fh:
        assert fh.read(5) == b"%PDF-", n

# ------------------------------------------------------------------- carta
CARTA = [
    "Dear Editors,",

    "Please consider the enclosed manuscript, “%s”, for publication in the Journal of "
    "Archaeological Method and Theory as an Original Article. It is original, is not under "
    "consideration elsewhere, and all data and code are openly deposited with a persistent "
    "identifier (https://doi.org/%s)." % (titulo, DOI_ESTE),

    "The paper addresses a gap between two things the discipline now rightly demands, and that "
    "this journal has done much to establish: computational reproducibility and computational "
    "correctness. A reproducible pipeline reproduces its errors faithfully. The manuscript "
    "characterises a class of defects in visibility analysis that raise no error, produce "
    "plausible output, survive both peer review and re-execution, and invert or erase the "
    "conclusion of the study containing them. Three were documented, and corrected, in an "
    "intervisibility study of the Titicaca basin whose pre-correction runs are openly deposited; "
    "that study is reported in a separate manuscript, and the present paper uses its "
    "three defects only as a specimen.",

    "Three methodological contributions follow. The first is an executable benchmark for "
    "line-of-sight engines, with %d terrain cases whose expectations are derived rather than "
    "written by hand and whose adequacy is measured by mutation analysis; that measurement caught "
    "the benchmark's own first set of behavioural properties detecting none of the injected "
    "defects, and the paper reports it as a finding. The second is a synthetic-landscape "
    "laboratory in which the truth is known by construction, so that what each design defect "
    "does to statistical inference is measured against an exact expectation. The third is a "
    "measurement of two widely used, freely available engines, gdal_viewshed (GDAL) and "
    "r.viewshed (GRASS), invoked directly with their shipped defaults. Both are correct when "
    "configured deliberately; at their defaults both fail, disagreeing on %d cases and returning "
    "the same wrong verdict on %d others. Agreement between two programs therefore validates "
    "nothing: only a benchmark with derived answers tells a correct engine from a correct engine "
    "badly configured. Two inexpensive diagnostics and a table of the parameters every visibility "
    "study should declare close the paper." % (n_casos, n_disc, n_comun),

    "The manuscript engages the journal's own conversation on reproducibility (Marwick, 2017) "
    "and on visibility networks (Brughmans, Keay, & Earl, 2015), and brings to it the testing "
    "literature of software engineering (metamorphic relations, mutation analysis and the oracle "
    "problem), which, to the author's knowledge, has not previously been applied to an "
    "archaeological geometric computation.",

    "The manuscript runs to about %s words of main text, plus %s words in tables, with %d figures, "
    "%d tables and %d references; the abstract has %d words. The figures are also supplied as "
    "separate files (PDF and TIFF). It is a single-author work; the author declares no competing "
    "interests and no funding. In accordance with the journal's policy, the use of a large "
    "language model as a coding and drafting assistant is documented in the Methods section, and "
    "the author takes full responsibility for the content."
    % ("{:,}".format(n_cuerpo), "{:,}".format(n_tablas), n_fig, n_tab, n_refs, n_abs),

    "Yours faithfully,",
    "%s\n%s\n%s, ORCID %s" % (AUTOR, FILIACION, CORREO, ORCID),
]
CARTA = [c.replace("'", "’") for c in CARTA]
assert "—" not in "".join(CARTA)

with open(os.path.join(OUT, "2_COVER_LETTER.txt"), "w", encoding="utf-8") as fh:
    fh.write("\n\n".join(CARTA) + "\n")

doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Cm(2.5)
sec.left_margin = sec.right_margin = Cm(2.5)
st = doc.styles["Normal"]
st.font.name = "Arial"
st.font.size = Pt(10.5)
st.paragraph_format.space_after = Pt(8)
st.paragraph_format.line_spacing = 1.15
for i, par in enumerate(CARTA):
    for linea in par.split("\n"):
        p = doc.add_paragraph()
        p.alignment = (WD_ALIGN_PARAGRAPH.LEFT if i in (0, len(CARTA) - 2, len(CARTA) - 1)
                       else WD_ALIGN_PARAGRAPH.JUSTIFY)
        r = p.add_run(linea)
        r.font.name = "Arial"
        r.font.size = Pt(10.5)
        if par.startswith("Yours") or i == len(CARTA) - 1:
            p.paragraph_format.space_after = Pt(2)
out_docx = os.path.join(OUT, "2_COVER_LETTER.docx")
doc.save(out_docx)
docmeta.limpiar(out_docx, autor=AUTOR, titulo="Cover letter - Journal of Archaeological Method and Theory")

print("Paquete -> %s" % OUT)
print("  manuscrito: %s palabras de cuerpo + %s en tablas | resumen %d | figuras %d | tablas %d | "
      "referencias %d" % ("{:,}".format(n_cuerpo), "{:,}".format(n_tablas), n_abs, n_fig, n_tab, n_refs))
print("  motores de fabrica: discrepan en %d, error comun en %d (de %d casos)" % (n_disc, n_comun, n_casos))
print("  ficheros:", sorted(os.listdir(OUT)))
