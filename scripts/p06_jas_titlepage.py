# -*- coding: utf-8 -*-
"""
Portada para Journal of Archaeological Science (revision de anonimo simple).

La portada va en fichero aparte del manuscrito y es la unica pieza con datos de
autoria: autor, filiacion, correspondencia, CRediT, conflictos, financiacion y
disponibilidad de datos. Elsevier exige CRediT por taxonomia; con autoria unica
se declaran todos los roles pertinentes.

Salida: ENVIO_JAS/3_TITLE_PAGE.docx
"""
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docmeta

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(BASE, "ENVIO_JAS", "3_TITLE_PAGE.docx")

AUTOR = "Milton Vladimir Mamani Calisaya"
FILIACION = "Universidad Nacional del Altiplano, Puno, Perú"
CORREO = "mmamanic@unap.edu.pe"
ORCID = "0000-0002-0676-0989"
DOI_CASO = "10.5281/zenodo.22176260"

doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Cm(2.5)
sec.left_margin = sec.right_margin = Cm(2.5)
st = doc.styles["Normal"]
st.font.name = "Arial"
st.font.size = Pt(10)
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.15


def P(t, size=10, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT,
      before=0, after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(t)
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p


P("Title page", 9, italic=True, after=14)

P("Reproducible and wrong: silent failures in archaeological visibility "
  "analysis and a benchmark to catch them", 14, True,
  align=WD_ALIGN_PARAGRAPH.CENTER, after=14)

P(AUTOR + " ¹ *", 11, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
P("¹ " + FILIACION, 10, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
P("* Corresponding author: %s — ORCID: https://orcid.org/%s" % (CORREO, ORCID),
  10, align=WD_ALIGN_PARAGRAPH.CENTER, after=16)

P("CRediT authorship contribution statement", 10.5, True, after=3)
P("%s: Conceptualization, Methodology, Software, Validation, Formal analysis, "
  "Investigation, Data curation, Writing – original draft, Writing – review & "
  "editing, Visualization." % AUTOR, after=10)

P("Declaration of competing interests", 10.5, True, after=3)
P("The author declares no competing financial interests or personal "
  "relationships that could have influenced the work reported in this paper.",
  after=10)

P("Funding", 10.5, True, after=3)
P("This research received no specific grant from any funding agency in the "
  "public, commercial, or not-for-profit sectors. It uses exclusively open "
  "data and open-source software.", after=10)

P("Data availability", 10.5, True, after=3)
P("The defect-injectable line-of-sight engine, the benchmark, the mutation "
  "harness, the landscape laboratory, all result files and the scripts that "
  "generate every figure and the manuscript itself are openly deposited under "
  "an MIT licence; the deposit DOI is cited in the manuscript's Data and code "
  "availability section. The field study whose defects are reproduced is "
  "likewise openly deposited, including its pre-correction runs "
  "(doi:%s)." % DOI_CASO, after=10)

P("Acknowledgements", 10.5, True, after=3)
P("None.", after=6)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
docmeta.limpiar(OUT, autor=AUTOR,
                titulo="Title page — Reproducible and wrong")
print("Portada -> %s" % OUT)
print("  metadatos:", docmeta.informe(OUT))
